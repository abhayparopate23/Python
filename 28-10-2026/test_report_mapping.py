"""Synthetic-workbook regression tests for report_mapping.py."""

from __future__ import annotations

from copy import copy
from pathlib import Path
import sys
import tempfile
import unittest

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill


sys.path.insert(0, str(Path(__file__).parent))
import report_mapping


BASE_RECORD = {
    "Email": "person@example.com",
    "Link": "https://www.linkedin.com/in/person",
    "QA Status": "Qualified",
    "Primary Reason": "",
    "Secondary Reason": "",
    "QA Comments": "",
    "First Experience": "Analyst\nAcme\n2022 - Present",
    "Job Title (As Per Experience)": "Analyst",
    "Company (As Per Experience)": "Acme",
    "Tenure (Overall)": "4 yrs",
    "Tenure (Current)": "2 yrs",
    "Experience's Geo": "Boston",
    "Geo (Header)": "New York",
    "Connections": "500+",
    "Followers": "120",
    "Activity": "6mo",
    "ID": "28-10-26-10-20-30",
}


def make_source(path: Path, records=None, headers=None):
    wb = Workbook()
    ws = wb.active
    ws.title = "Profiles"
    headers = headers or report_mapping.MAPPED_HEADERS
    ws.append(headers)
    for record in records if records is not None else [BASE_RECORD]:
        ws.append([record.get(header) for header in headers])
    wb.save(path)
    wb.close()


def make_report(path: Path, rows=None, mapping_headers=False):
    wb = Workbook()
    ws = wb.active
    ws.title = "Report"
    ws.cell(1, 1, "Email ID")
    ws.cell(1, 2, "Job Title")
    ws.cell(1, 3, "Company Name")
    ws.cell(1, 53, "Original BA")
    if mapping_headers:
        for offset, header in enumerate(report_mapping.MAPPED_HEADERS):
            ws.cell(1, 54 + offset, header)
    for row_num, row in enumerate(rows if rows is not None else [
        {"Email ID": "person@example.com", "Job Title": "Analyst", "Company Name": "Acme"}
    ], start=2):
        ws.cell(row_num, 1, row.get("Email ID"))
        ws.cell(row_num, 2, row.get("Job Title"))
        ws.cell(row_num, 3, row.get("Company Name"))
        ws.cell(row_num, 53, row.get("BA", f"keep-{row_num}"))
    wb.save(path)
    wb.close()


class ReportMappingTests(unittest.TestCase):
    def setUp(self):
        # The managed Windows runner restricts its global temp folder; keep all
        # synthetic workbook I/O inside this authorized project directory.
        self.temp = tempfile.TemporaryDirectory(dir=Path(__file__).parent)
        self.root = Path(self.temp.name)
        self.source = self.root / "generated.xlsx"
        self.report = self.root / "report.xlsx"

    def tearDown(self):
        self.temp.cleanup()

    def run_mapping(self, source_records=None, report_rows=None, mapping_headers=False):
        make_source(self.source, source_records)
        make_report(self.report, report_rows, mapping_headers=mapping_headers)
        return report_mapping.process_files(self.source, self.report)

    def mapped(self, row=2, column=54):
        wb = load_workbook(self.report)
        value = wb["Report"].cell(row, column).value
        wb.close()
        return value

    def test_complete_exact_email_match_maps_all_fields(self):
        summary = self.run_mapping()
        self.assertEqual(summary.matched_rows, 1)
        self.assertEqual(self.mapped(2, 54), BASE_RECORD["Email"])
        self.assertEqual(self.mapped(2, 70), BASE_RECORD["ID"])

    def test_email_match_trims_and_ignores_case(self):
        record = dict(BASE_RECORD, Email="  Person@Example.com ")
        summary = self.run_mapping([record], [{"Email ID": "person@example.COM ", "Job Title": "Analyst", "Company Name": "Acme"}])
        self.assertEqual(summary.matched_rows, 1)
        self.assertEqual(self.mapped(), "  Person@Example.com ")
        self.assertEqual(len(summary.mismatches["Email"]), 0)

    def test_job_title_mismatch_highlights_only_job_title_cell(self):
        summary = self.run_mapping(report_rows=[{"Email ID": BASE_RECORD["Email"], "Job Title": "Director", "Company Name": "Acme"}])
        self.assertEqual(len(summary.mismatches["Job Title (As Per Experience)"]), 1)
        wb = load_workbook(self.report)
        ws = wb["Report"]
        self.assertEqual(ws["BI2"].font.color.rgb, report_mapping.RED)
        self.assertFalse(report_mapping._is_red(ws["BJ2"].font.color))
        self.assertFalse(report_mapping._is_red(ws["BB2"].font.color))
        wb.close()

    def test_company_mismatch_highlights_only_company_cell(self):
        summary = self.run_mapping(report_rows=[{"Email ID": BASE_RECORD["Email"], "Job Title": "Analyst", "Company Name": "Other Co"}])
        self.assertEqual(len(summary.mismatches["Company (As Per Experience)"]), 1)
        wb = load_workbook(self.report)
        ws = wb["Report"]
        self.assertEqual(ws["BJ2"].font.color.rgb, report_mapping.RED)
        self.assertFalse(report_mapping._is_red(ws["BI2"].font.color))
        wb.close()

    def test_different_email_cannot_be_a_mapped_email_mismatch(self):
        # Email is the join key, so a genuinely different email must be unmatched.
        summary = self.run_mapping(report_rows=[{"Email ID": "other@example.com", "Job Title": "Analyst", "Company Name": "Acme"}])
        self.assertEqual(summary.matched_rows, 0)
        self.assertEqual(len(summary.unmatched_rows), 1)
        self.assertIsNone(self.mapped())
        self.assertFalse(report_mapping.normalize_comparison("other@example.com") == report_mapping.normalize_comparison(BASE_RECORD["Email"]))

    def test_unmatched_report_email_is_blank(self):
        summary = self.run_mapping(report_rows=[{"Email ID": "missing@example.com", "Job Title": "Analyst", "Company Name": "Acme"}])
        self.assertEqual(len(summary.unmatched_rows), 1)
        self.assertEqual(self.mapped(), None)

    def test_duplicate_generated_email_is_ambiguous_and_blank(self):
        records = [dict(BASE_RECORD), dict(BASE_RECORD, Link="https://www.linkedin.com/in/other")]
        summary = self.run_mapping(records)
        self.assertEqual(summary.matched_rows, 0)
        self.assertEqual(summary.ambiguous_generated, [(2, BASE_RECORD["Email"], 2)])
        self.assertIsNone(self.mapped())

    def test_duplicate_report_email_rows_are_not_chosen_arbitrarily(self):
        rows = [
            {"Email ID": BASE_RECORD["Email"], "Job Title": "Analyst", "Company Name": "Acme"},
            {"Email ID": BASE_RECORD["Email"], "Job Title": "Analyst", "Company Name": "Acme"},
        ]
        summary = self.run_mapping(report_rows=rows)
        self.assertEqual(summary.duplicate_report_keys, [(BASE_RECORD["Email"], [2, 3])])
        self.assertEqual(summary.matched_rows, 0)
        self.assertIsNone(self.mapped(2))
        self.assertIsNone(self.mapped(3))

    def test_blank_report_email_does_not_match_blank_source_email(self):
        record = dict(BASE_RECORD, Email="")
        summary = self.run_mapping([record], [{"Email ID": "", "Job Title": "Analyst", "Company Name": "Acme"}])
        self.assertEqual(summary.missing_email_rows, [2])
        self.assertEqual(summary.matched_rows, 0)
        self.assertIsNone(self.mapped())

    def test_missing_comparison_values_are_counted_not_red(self):
        rows = [{"Email ID": BASE_RECORD["Email"], "Job Title": "", "Company Name": "Acme"}]
        summary = self.run_mapping(report_rows=rows)
        self.assertIn((2, BASE_RECORD["Email"], "Job Title (As Per Experience)"), summary.missing_comparison_values)
        self.assertEqual(len(summary.mismatches["Job Title (As Per Experience)"]), 0)
        wb = load_workbook(self.report)
        self.assertFalse(report_mapping._is_red(wb["Report"]["BI2"].font.color))
        wb.close()

    def test_existing_destination_headers_in_bb_br_are_updated_in_place(self):
        summary = self.run_mapping(mapping_headers=True)
        self.assertEqual(summary.matched_rows, 1)
        wb = load_workbook(self.report)
        ws = wb["Report"]
        self.assertEqual([ws.cell(1, c).value for c in range(54, 71)], report_mapping.MAPPED_HEADERS)
        wb.close()

    def test_conflicting_destination_headers_stop_without_modifying_report(self):
        make_source(self.source)
        make_report(self.report)
        wb = load_workbook(self.report)
        wb["Report"]["BB1"] = "Unrelated field"
        wb.save(self.report)
        wb.close()
        before = self.report.read_bytes()
        with self.assertRaises(report_mapping.MappingError):
            report_mapping.process_files(self.source, self.report)
        self.assertEqual(self.report.read_bytes(), before)

    def test_columns_through_ba_remain_unchanged(self):
        make_source(self.source)
        make_report(self.report)
        before_wb = load_workbook(self.report, data_only=False)
        before = [[before_wb["Report"].cell(r, c).value for c in range(1, 54)] for r in range(1, 3)]
        before_wb.close()
        report_mapping.process_files(self.source, self.report)
        after_wb = load_workbook(self.report, data_only=False)
        after = [[after_wb["Report"].cell(r, c).value for c in range(1, 54)] for r in range(1, 3)]
        after_wb.close()
        self.assertEqual(after, before)

    def test_matching_cells_retain_existing_font_and_fill_formatting(self):
        make_source(self.source)
        make_report(self.report, mapping_headers=True)
        wb = load_workbook(self.report)
        cell = wb["Report"]["BB2"]
        cell.font = Font(name="Calibri", size=13, bold=True, italic=True, color="FF123456")
        cell.fill = PatternFill("solid", fgColor="FFFFFF00")
        wb.save(self.report)
        wb.close()
        report_mapping.process_files(self.source, self.report)
        wb = load_workbook(self.report)
        cell = wb["Report"]["BB2"]
        self.assertEqual((cell.font.name, cell.font.sz, cell.font.bold, cell.font.italic, cell.font.color.rgb), ("Calibri", 13, True, True, "FF123456"))
        self.assertEqual(cell.fill.fgColor.rgb, "FFFFFF00")
        wb.close()

    def test_only_mismatched_cells_get_red_font(self):
        rows = [{"Email ID": BASE_RECORD["Email"], "Job Title": "Different", "Company Name": "Other"}]
        summary = self.run_mapping(report_rows=rows)
        self.assertEqual(summary.red_cells, 2)
        wb = load_workbook(self.report)
        ws = wb["Report"]
        self.assertFalse(report_mapping._is_red(ws["BB2"].font.color))
        self.assertEqual(ws["BI2"].font.color.rgb, report_mapping.RED)
        self.assertEqual(ws["BJ2"].font.color.rgb, report_mapping.RED)
        self.assertFalse(report_mapping._is_red(ws["BH2"].font.color))
        wb.close()

    def test_stale_red_is_cleared_on_later_matching_run(self):
        rows = [{"Email ID": BASE_RECORD["Email"], "Job Title": "Different", "Company Name": "Acme"}]
        self.run_mapping(report_rows=rows)
        wb = load_workbook(self.report)
        ws = wb["Report"]
        self.assertTrue(report_mapping._is_red(ws["BI2"].font.color))
        font = copy(ws["BA2"].font)
        font.color = "FF0055AA"
        ws["BA2"].font = font
        ws["B2"] = "Analyst"
        wb.save(self.report)
        wb.close()
        report_mapping.process_files(self.source, self.report)
        wb = load_workbook(self.report)
        self.assertEqual(wb["Report"]["BI2"].font.color.rgb, "FF0055AA")
        wb.close()

    def test_repeated_execution_does_not_duplicate_headers(self):
        make_source(self.source)
        make_report(self.report)
        report_mapping.process_files(self.source, self.report)
        report_mapping.process_files(self.source, self.report)
        wb = load_workbook(self.report)
        ws = wb["Report"]
        self.assertEqual([ws.cell(1, c).value for c in range(54, 71)], report_mapping.MAPPED_HEADERS)
        wb.close()

    def test_validation_failure_leaves_report_unchanged(self):
        self.source.write_text("not an xlsx", encoding="utf-8")
        make_report(self.report)
        before = self.report.read_bytes()
        with self.assertRaises(report_mapping.MappingError):
            report_mapping.process_files(self.source, self.report)
        self.assertEqual(self.report.read_bytes(), before)
        self.assertEqual(list(self.root.glob("Report_backup_*.xlsx")), [])

    def test_backup_is_created_before_report_is_replaced(self):
        make_source(self.source)
        make_report(self.report)
        summary = report_mapping.process_files(self.source, self.report)
        backup = Path(summary.backup_path)
        self.assertTrue(backup.exists())
        wb = load_workbook(backup, read_only=True)
        self.assertIsNone(wb["Report"].cell(1, 54).value)
        wb.close()
        self.assertEqual(self.mapped(), BASE_RECORD["Email"])

    def test_same_file_is_refused_without_changes(self):
        make_source(self.source)
        before = self.source.read_bytes()
        with self.assertRaisesRegex(report_mapping.MappingError, "same file"):
            report_mapping.process_files(self.source, self.source)
        self.assertEqual(self.source.read_bytes(), before)

    def test_report_header_row_can_be_below_title_rows(self):
        make_source(self.source)
        wb = Workbook()
        ws = wb.active
        ws.title = "Report"
        ws["A1"] = "Quarterly Report"
        ws["A3"] = "Email ID"
        ws["B3"] = "Job Title"
        ws["C3"] = "Company Name"
        ws["BA3"] = "Original BA"
        ws["A4"] = BASE_RECORD["Email"]
        ws["B4"] = "Analyst"
        ws["C4"] = "Acme"
        ws["BA4"] = "keep"
        wb.save(self.report)
        wb.close()
        summary = report_mapping.process_files(self.source, self.report)
        self.assertEqual(summary.matched_rows, 1)
        self.assertEqual(self.mapped(4, 54), BASE_RECORD["Email"])

    def test_ambiguous_report_header_locations_are_rejected(self):
        make_source(self.source)
        make_report(self.report)
        wb = load_workbook(self.report)
        ws = wb["Report"]
        ws["A5"] = "Email ID"
        ws["B5"] = "Job Title"
        ws["C5"] = "Company Name"
        wb.save(self.report)
        wb.close()
        before = self.report.read_bytes()
        with self.assertRaisesRegex(report_mapping.MappingError, "More than one Report header row"):
            report_mapping.process_files(self.source, self.report)
        self.assertEqual(self.report.read_bytes(), before)


if __name__ == "__main__":
    unittest.main(verbosity=2)
