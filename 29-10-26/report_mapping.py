"""Map LinkedIn Experience Copier records into a Report workbook.

Run with ``python report_mapping.py`` after editing the two paths below.
The Report workbook is updated in place after a timestamped backup is made.
"""

from __future__ import annotations

from collections import Counter, defaultdict
from copy import copy
from dataclasses import dataclass, field
from datetime import datetime
import os
from pathlib import Path
import re
import shutil
import sys
import tempfile
import unicodedata
import warnings
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from openpyxl.styles import Color
from openpyxl.utils import get_column_letter


# ============================================================
# REPORT MAPPING CONFIGURATION
# Edit these paths before running the script.
# ============================================================

GENERATED_FILE_PATH = r"C:\CHANGE\THIS\TO\GENERATED_FILE.xlsx"
REPORT_FILE_PATH = r"C:\CHANGE\THIS\TO\REPORT.xlsx"


MAPPED_HEADERS = [
    "Email",
    "Link",
    "QA Status",
    "Primary Reason",
    "Secondary Reason",
    "QA Comments",
    "First Experience",
    "Job Title (As Per Experience)",
    "Company (As Per Experience)",
    "Tenure (Overall)",
    "Tenure (Current)",
    "Experience's Geo",
    "Geo (Header)",
    "Connections",
    "Followers",
    "Activity",
    "ID",
]

# Destination field -> accepted source-header spellings. Canonical names are
# tried first; aliases exist only for known earlier generated-workbook schemas.
SOURCE_HEADER_ALIASES = {
    "Email": ("Email",),
    "Link": ("Link", "LinkedIn URL", "LinkedIn Link"),
    "QA Status": ("QA Status",),
    "Primary Reason": ("Primary Reason",),
    "Secondary Reason": ("Secondary Reason",),
    "QA Comments": ("QA Comments",),
    "First Experience": ("First Experience",),
    "Job Title (As Per Experience)": (
        "Job Title (As Per Experience)", "Job Title", "First Experience Job Title"
    ),
    "Company (As Per Experience)": (
        "Company (As Per Experience)", "Company", "First Experience Company"
    ),
    "Tenure (Overall)": ("Tenure (Overall)", "Overall Tenure"),
    "Tenure (Current)": ("Tenure (Current)", "Current Tenure"),
    "Experience's Geo": ("Experience's Geo", "Experience Geo", "First Experience Geo"),
    "Geo (Header)": ("Geo (Header)", "Geo", "Header Geo"),
    "Connections": ("Connections",),
    "Followers": ("Followers",),
    "Activity": ("Activity",),
    "ID": ("ID", "Profile ID"),
}

COMPARISON_FIELDS = {
    "Email": "Email ID",
    "Job Title (As Per Experience)": "Job Title",
    "Company (As Per Experience)": "Company Name",
}
REPORT_HEADER_ROW_SCAN = 50
SOURCE_HEADER_ROW_SCAN = 10
FIRST_MAPPING_COLUMN = 54  # BB; BA is 53.
LAST_MAPPING_COLUMN = FIRST_MAPPING_COLUMN + len(MAPPED_HEADERS) - 1  # BR.
RED = "FFFF0000"


class MappingError(RuntimeError):
    """A validation or workbook-layout error that must stop the mapping."""


@dataclass
class SourceRecord:
    values: dict[str, object]
    sheet_name: str
    row_number: int


@dataclass
class MappingSummary:
    generated_path: str
    report_path: str
    backup_path: str = ""
    report_rows_processed: int = 0
    matched_rows: int = 0
    unmatched_rows: list[tuple[int, str]] = field(default_factory=list)
    ambiguous_generated: list[tuple[int, str, int]] = field(default_factory=list)
    duplicate_report_keys: list[tuple[str, list[int]]] = field(default_factory=list)
    missing_email_rows: list[int] = field(default_factory=list)
    mismatches: dict[str, list[tuple[int, str]]] = field(
        default_factory=lambda: {name: [] for name in COMPARISON_FIELDS}
    )
    missing_comparison_values: list[tuple[int, str, str]] = field(default_factory=list)
    red_cells: int = 0


def normalize_header(value: object) -> str:
    """Normalize a header by ignoring case, accents, spacing, and punctuation."""
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = "".join(char for char in text if not unicodedata.combining(char))
    return "".join(char.casefold() for char in text if char.isalnum())


def normalize_email(value: object) -> str:
    """Return the email join key; blank values deliberately have no key."""
    if value is None:
        return ""
    return str(value).strip().casefold()


def normalize_comparison(value: object) -> str:
    """Normalize text for exact, case-insensitive comparison (never fuzzy)."""
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value).strip()).casefold()


def column_header_map(values: tuple[object, ...] | list[object]) -> dict[str, int]:
    """Resolve canonical source fields to 1-based columns, rejecting ambiguity."""
    normalized = defaultdict(list)
    for column, value in enumerate(values, start=1):
        key = normalize_header(value)
        if key:
            normalized[key].append(column)

    resolved: dict[str, int] = {}
    ambiguous: list[str] = []
    for destination, aliases in SOURCE_HEADER_ALIASES.items():
        canonical_key = normalize_header(destination)
        exact = normalized.get(canonical_key, [])
        if len(exact) > 1:
            ambiguous.append(destination)
            continue
        if len(exact) == 1:
            resolved[destination] = exact[0]
            continue
        alias_matches = []
        for alias in aliases[1:]:
            alias_matches.extend(normalized.get(normalize_header(alias), []))
        alias_matches = sorted(set(alias_matches))
        if len(alias_matches) > 1:
            ambiguous.append(destination)
        elif alias_matches:
            resolved[destination] = alias_matches[0]

    if ambiguous:
        raise MappingError("Ambiguous source headers for: " + ", ".join(ambiguous))
    return resolved


def _find_source_header(ws):
    observed_missing: set[str] = set(MAPPED_HEADERS)
    for row_number in range(1, min(ws.max_row, SOURCE_HEADER_ROW_SCAN) + 1):
        values = [ws.cell(row_number, col).value for col in range(1, ws.max_column + 1)]
        try:
            headers = column_header_map(values)
        except MappingError:
            continue
        if "Email" in headers and "Link" in headers:
            observed_missing &= set(MAPPED_HEADERS) - set(headers)
            if len(headers) == len(MAPPED_HEADERS):
                return row_number, headers
    return None, observed_missing


def _read_source_records(workbook) -> list[SourceRecord]:
    records: list[SourceRecord] = []
    source_sheets = []
    missing_by_sheet = []
    for ws in workbook.worksheets:
        header_row, result = _find_source_header(ws)
        if header_row is None:
            missing = sorted(result)
            if missing != sorted(MAPPED_HEADERS):
                missing_by_sheet.append((ws.title, missing))
            continue
        headers = result
        source_sheets.append(ws.title)
        for row_number in range(header_row + 1, ws.max_row + 1):
            if not any(ws.cell(row_number, col).value is not None for col in headers.values()):
                continue
            values = {
                field_name: ws.cell(row_number, source_column).value
                for field_name, source_column in headers.items()
            }
            records.append(SourceRecord(values, ws.title, row_number))
    if not source_sheets:
        details = []
        if missing_by_sheet:
            for sheet, missing in missing_by_sheet:
                details.append(f"{sheet}: missing {', '.join(missing)}")
        detail = "\n".join(details) or "No worksheet with Email and Link headers was found."
        raise MappingError("Could not identify a generated-data header row. " + detail)
    return records


def _resolve_report_header(value: object) -> str | None:
    key = normalize_header(value)
    aliases = {
        "emailid": "Email ID",
        "jobtitle": "Job Title",
        "companyname": "Company Name",
    }
    return aliases.get(key)


def _find_report_layout(workbook):
    candidates = []
    for ws in workbook.worksheets:
        for row_number in range(1, min(ws.max_row, REPORT_HEADER_ROW_SCAN) + 1):
            found: dict[str, list[int]] = defaultdict(list)
            for col in range(1, ws.max_column + 1):
                label = _resolve_report_header(ws.cell(row_number, col).value)
                if label:
                    found[label].append(col)
            required = set(COMPARISON_FIELDS.values())
            if required.issubset(found):
                duplicates = [label for label, cols in found.items() if len(cols) > 1]
                if duplicates:
                    raise MappingError(
                        f"Report sheet {ws.title!r}, row {row_number}, has duplicate headers: "
                        + ", ".join(duplicates)
                    )
                candidates.append((ws, row_number, {name: cols[0] for name, cols in found.items()}))
    if not candidates:
        raise MappingError(
            "Could not find one Report header row containing Email ID, Job Title, and Company Name. "
            "The Report workbook is not included in the project, so verify those labels and their row."
        )
    if len(candidates) != 1:
        locations = ", ".join(f"{ws.title}!row {row}" for ws, row, _ in candidates)
        raise MappingError("More than one Report header row matches the required fields: " + locations)
    return candidates[0]


def _validate_destination_layout(ws, header_row: int) -> None:
    expected = [normalize_header(title) for title in MAPPED_HEADERS]
    target_values = [
        normalize_header(ws.cell(header_row, col).value)
        for col in range(FIRST_MAPPING_COLUMN, LAST_MAPPING_COLUMN + 1)
    ]
    if any(target_values) and target_values != expected:
        actual = [ws.cell(header_row, col).value for col in range(FIRST_MAPPING_COLUMN, LAST_MAPPING_COLUMN + 1)]
        raise MappingError(
            f"Report {ws.title!r} has conflicting headers in BB:BR: {actual!r}. "
            "Expected the 17 mapping headers in the requested order. No Report changes were made."
        )

    mapping_titles = set(expected)
    misplaced = []
    for col in range(1, ws.max_column + 1):
        if FIRST_MAPPING_COLUMN <= col <= LAST_MAPPING_COLUMN:
            continue
        if normalize_header(ws.cell(header_row, col).value) in mapping_titles:
            misplaced.append(f"{get_column_letter(col)}:{ws.cell(header_row, col).value}")
    if misplaced:
        raise MappingError(
            "Mapping headers already exist outside BB:BR: " + ", ".join(misplaced)
            + ". Resolve the layout conflict before running. No Report changes were made."
        )

    if target_values != expected:
        for row in range(1, ws.max_row + 1):
            if row == header_row:
                continue
            if any(
                ws.cell(row, col).value is not None
                for col in range(FIRST_MAPPING_COLUMN, LAST_MAPPING_COLUMN + 1)
            ):
                raise MappingError(
                    f"Report contains data in BB:BR without the expected mapping headers "
                    f"(first found at row {row}). No Report changes were made."
                )


def _unsupported_parts(path: Path) -> list[str]:
    unsupported = []
    with ZipFile(path) as archive:
        names = set(archive.namelist())
    known_prefixes = (
        "xl/activeX/", "xl/ctrlProps/", "xl/slicerCaches/", "xl/slicers/",
        "xl/pivotCache/", "xl/pivotTables/", "xl/threadedComments/", "customUI/",
    )
    for prefix in known_prefixes:
        if any(name.startswith(prefix) for name in names):
            unsupported.append(prefix.rstrip("/"))
    if any(name.lower().endswith("vbaproject.bin") for name in names):
        unsupported.append("VBA macros")
    return unsupported


def _load_validated_workbook(path: Path, label: str):
    try:
        with warnings.catch_warnings(record=True) as caught:
            warnings.simplefilter("always")
            workbook = load_workbook(path, data_only=False, read_only=False, keep_links=True)
        extension_warnings = [
            str(item.message) for item in caught
            if "extension" in str(item.message).casefold() or "unsupported" in str(item.message).casefold()
        ]
        if extension_warnings:
            workbook.close()
            raise MappingError(
                f"{label} workbook uses features openpyxl reports as unsupported: "
                + "; ".join(extension_warnings)
            )
        unsupported = _unsupported_parts(path)
        if unsupported:
            workbook.close()
            raise MappingError(
                f"{label} workbook contains features this script will not risk changing: "
                + ", ".join(unsupported)
            )
        return workbook
    except (BadZipFile, OSError, ValueError, KeyError) as exc:
        if isinstance(exc, MappingError):
            raise
        raise MappingError(f"{label} is not a readable supported Excel workbook: {exc}") from exc


def _unique_backup_path(report_path: Path) -> Path:
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    base = report_path.parent / f"Report_backup_{stamp}"
    candidate = base.with_suffix(report_path.suffix)
    suffix = 2
    while candidate.exists():
        candidate = report_path.parent / f"Report_backup_{stamp}_{suffix:02d}{report_path.suffix}"
        suffix += 1
    return candidate


def create_backup(report_path: Path) -> Path:
    backup = _unique_backup_path(report_path)
    try:
        # Exclusive creation prevents replacing a backup created by another run.
        with report_path.open("rb") as source, backup.open("xb") as destination:
            shutil.copyfileobj(source, destination)
        shutil.copystat(report_path, backup)
        return backup
    except OSError as exc:
        try:
            backup.unlink(missing_ok=True)
        except OSError:
            pass
        raise MappingError(f"Could not create Report backup {backup}: {exc}") from exc


def _is_red(color) -> bool:
    return bool(color and color.type == "rgb" and str(color.rgb).upper() in {RED, "00FF0000"})


def _set_font_color(cell, color) -> None:
    font = copy(cell.font)
    font.color = copy(color) if color is not None else None
    cell.font = font


def _copy_row_style_if_unstyled(cell, template_cell) -> None:
    if not cell.has_style and template_cell.has_style:
        cell._style = copy(template_cell._style)


def _write_headers(ws, header_row: int) -> None:
    template = ws.cell(header_row, FIRST_MAPPING_COLUMN - 1)
    for offset, title in enumerate(MAPPED_HEADERS):
        cell = ws.cell(header_row, FIRST_MAPPING_COLUMN + offset)
        if not cell.has_style and template.has_style:
            cell._style = copy(template._style)
        cell.value = title


def _report_data_rows(ws, header_row: int) -> list[int]:
    return [
        row for row in range(header_row + 1, ws.max_row + 1)
        if any(ws.cell(row, col).value is not None for col in range(1, 54))
    ]


def _apply_mapping(source_records, report_ws, header_row, report_headers, summary):
    generated_by_email: dict[str, list[SourceRecord]] = defaultdict(list)
    for record in source_records:
        key = normalize_email(record.values.get("Email"))
        if key:
            generated_by_email[key].append(record)

    report_rows = _report_data_rows(report_ws, header_row)
    summary.report_rows_processed = len(report_rows)
    report_rows_by_email: dict[str, list[int]] = defaultdict(list)
    for row in report_rows:
        key = normalize_email(report_ws.cell(row, report_headers["Email ID"]).value)
        if key:
            report_rows_by_email[key].append(row)
        else:
            summary.missing_email_rows.append(row)
    summary.duplicate_report_keys = sorted(
        [(key, rows) for key, rows in report_rows_by_email.items() if len(rows) > 1],
        key=lambda item: item[0],
    )
    duplicate_rows = {row for _, rows in summary.duplicate_report_keys for row in rows}

    _write_headers(report_ws, header_row)
    destination_column = {title: FIRST_MAPPING_COLUMN + offset for offset, title in enumerate(MAPPED_HEADERS)}
    source_comparison = {
        "Email": "Email ID",
        "Job Title (As Per Experience)": "Job Title",
        "Company (As Per Experience)": "Company Name",
    }

    for row in report_rows:
        report_email = report_ws.cell(row, report_headers["Email ID"]).value
        key = normalize_email(report_email)
        candidates = generated_by_email.get(key, []) if key else []
        # Never leave stale mappings or script-applied red from a previous run.
        for field_name, comparison_header in source_comparison.items():
            mapped_cell = report_ws.cell(row, destination_column[field_name])
            if _is_red(mapped_cell.font.color):
                # New mapping cells inherit the Report row's existing BA style;
                # restore that row's base font color after a previous red mismatch.
                _set_font_color(mapped_cell, report_ws.cell(row, FIRST_MAPPING_COLUMN - 1).font.color)
        for title, col in destination_column.items():
            target = report_ws.cell(row, col)
            _copy_row_style_if_unstyled(target, report_ws.cell(row, FIRST_MAPPING_COLUMN - 1))
            target.value = None

        if not key:
            continue
        if row in duplicate_rows:
            continue
        if not candidates:
            summary.unmatched_rows.append((row, str(report_email).strip()))
            continue
        if len(candidates) > 1:
            summary.ambiguous_generated.append((row, str(report_email).strip(), len(candidates)))
            continue

        record = candidates[0]
        summary.matched_rows += 1
        for field_name, col in destination_column.items():
            report_ws.cell(row, col).value = record.values.get(field_name)

        for field_name, comparison_header in COMPARISON_FIELDS.items():
            mapped_cell = report_ws.cell(row, destination_column[field_name])
            original_cell = report_ws.cell(row, report_headers[comparison_header])
            mapped_value = mapped_cell.value
            original_value = original_cell.value
            if not normalize_comparison(mapped_value) or not normalize_comparison(original_value):
                summary.missing_comparison_values.append(
                    (row, str(report_email).strip(), field_name)
                )
                # Missing data is not a mismatch; stale script red was cleared above.
                continue
            if normalize_comparison(mapped_value) != normalize_comparison(original_value):
                font = copy(mapped_cell.font)
                font.color = RED
                mapped_cell.font = font
                summary.mismatches[field_name].append((row, str(report_email).strip()))
                summary.red_cells += 1


def process_files(generated_file: str | Path, report_file: str | Path) -> MappingSummary:
    generated_path = Path(generated_file).expanduser()
    report_path = Path(report_file).expanduser()
    if not generated_path.exists() or not generated_path.is_file():
        raise MappingError(f"Generated file does not exist or is not a file: {generated_path}")
    if not report_path.exists() or not report_path.is_file():
        raise MappingError(f"Report file does not exist or is not a file: {report_path}")
    try:
        if os.path.samefile(generated_path, report_path):
            raise MappingError("Generated file and Report resolve to the same file. No changes were made.")
    except OSError as exc:
        raise MappingError(f"Could not compare input paths: {exc}") from exc

    generated_path = generated_path.resolve()
    report_path = report_path.resolve()
    generated_wb = report_wb = None
    try:
        generated_wb = _load_validated_workbook(generated_path, "Generated")
        report_wb = _load_validated_workbook(report_path, "Report")
        source_records = _read_source_records(generated_wb)
        report_ws, header_row, report_headers = _find_report_layout(report_wb)
        _validate_destination_layout(report_ws, header_row)

        summary = MappingSummary(str(generated_path), str(report_path))
        _apply_mapping(source_records, report_ws, header_row, report_headers, summary)

        # The original on-disk Report is untouched until all validation and
        # mapping operations above have completed and this backup has succeeded.
        backup_path = create_backup(report_path)
        summary.backup_path = str(backup_path)
        temp_path = None
        try:
            fd, temp_name = tempfile.mkstemp(
                prefix=f".{report_path.stem}.mapping-", suffix=report_path.suffix,
                dir=report_path.parent,
            )
            os.close(fd)
            temp_path = Path(temp_name)
            report_wb.save(temp_path)
            # Verify that the temporary workbook is readable before replacing Report.
            check_wb = load_workbook(temp_path, read_only=True, data_only=False)
            check_wb.close()
            os.replace(temp_path, report_path)
            temp_path = None
        except PermissionError as exc:
            raise MappingError(
                f"Could not replace the Report; it may be open or locked in Excel. "
                f"Close the workbook in Excel and rerun. The backup is preserved at {backup_path}."
            ) from exc
        except OSError as exc:
            raise MappingError(
                f"Could not safely save the Report: {exc}. The original Report remains in place; "
                f"the backup is preserved at {backup_path}."
            ) from exc
        finally:
            if temp_path is not None:
                try:
                    temp_path.unlink(missing_ok=True)
                except OSError:
                    pass
        return summary
    finally:
        if generated_wb is not None:
            generated_wb.close()
        if report_wb is not None:
            report_wb.close()


def print_summary(summary: MappingSummary) -> None:
    print("Report mapping completed.")
    print(f"Generated file: {summary.generated_path}")
    print(f"Report file: {summary.report_path}")
    print(f"Backup file: {summary.backup_path}")
    print(f"Report data rows processed: {summary.report_rows_processed}")
    print(f"Successfully matched rows: {summary.matched_rows}")
    print(f"Unmatched Report emails: {len(summary.unmatched_rows)}")
    print(f"Ambiguous generated-email matches: {len(summary.ambiguous_generated)}")
    print(f"Duplicate Report email keys: {len(summary.duplicate_report_keys)}")
    print(f"Rows with missing email keys: {len(summary.missing_email_rows)}")
    for field_name, label in (
        ("Email", "Email mismatches"),
        ("Job Title (As Per Experience)", "Job Title mismatches"),
        ("Company (As Per Experience)", "Company mismatches"),
    ):
        print(f"{label}: {len(summary.mismatches[field_name])}")
    print(f"Missing comparison values: {len(summary.missing_comparison_values)}")
    print(f"Cells highlighted red: {summary.red_cells}")
    if summary.unmatched_rows:
        print("Unmatched rows (Report row, email):")
        for row, email in summary.unmatched_rows:
            print(f"  row {row}: {email}")
    if summary.ambiguous_generated:
        print("Ambiguous generated matches (Report row, email, candidate count):")
        for row, email, count in summary.ambiguous_generated:
            print(f"  row {row}: {email} ({count} candidates)")
    if summary.duplicate_report_keys:
        print("Duplicate Report email keys (email, Report rows):")
        for email, rows in summary.duplicate_report_keys:
            print(f"  {email}: {', '.join(map(str, rows))}")
    if summary.missing_email_rows:
        print("Missing Report email keys (Report rows): " + ", ".join(map(str, summary.missing_email_rows)))
    if summary.missing_comparison_values:
        print("Missing comparison values (Report row, email, field):")
        for row, email, field_name in summary.missing_comparison_values:
            print(f"  row {row}: {email} — {field_name}")
    for field_name, issues in summary.mismatches.items():
        if issues:
            print(f"{field_name} mismatch rows (Report row, email):")
            for row, email in issues:
                print(f"  row {row}: {email}")
    print(f"Final saved Report path: {summary.report_path}")


def main() -> int:
    if "CHANGE" in GENERATED_FILE_PATH or "CHANGE" in REPORT_FILE_PATH:
        print(
            "Edit GENERATED_FILE_PATH and REPORT_FILE_PATH at the top of report_mapping.py "
            "before running it.",
            file=sys.stderr,
        )
        return 2
    try:
        summary = process_files(GENERATED_FILE_PATH, REPORT_FILE_PATH)
    except MappingError as exc:
        print(f"Report mapping stopped: {exc}", file=sys.stderr)
        return 1
    print_summary(summary)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
